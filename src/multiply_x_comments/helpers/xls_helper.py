from openpyxl import load_workbook


def read_xslsx_file(file_path: str) -> list[tuple] | None:
    values = []
    try:
        workbook = load_workbook(file_path)
        sheet = workbook.active
        for row in sheet.iter_rows(min_col=1, max_col=2, values_only=True):
            text = row[0]
            number = row[1]
            if text is not None:
                values.append(
                    (str(text), int(number),)
                )
    except FileNotFoundError:
        print(f"File not found: {file_path}")
    except Exception as e:
        print(f"An error occurred: {e}")
    return values    